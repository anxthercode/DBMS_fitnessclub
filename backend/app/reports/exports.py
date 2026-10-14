"""One sales dataset, three portable file formats. No spreadsheet runtime/service."""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from app.memberships.products import MINSK

TEXT = {
    "ru": {
        "title": "Отчёт о продажах NORTHSIDE",
        "period": "Период",
        "as_of": "Данные на",
        "note": "Успешные имитации оплаты. Даты по Минску, суммы в BYN после скидок.",
        "total": "Итого",
        "empty": "За выбранный период оплаченных заказов нет.",
        "headers": ["Заказ", "Дата оплаты", "Позиций", "До скидки", "Скидка", "Итого BYN"],
        "sheet": "Продажи",
        "source": "Источник: PostgreSQL NORTHSIDE, заказы и успешные платежи.",
    },
    "en": {
        "title": "NORTHSIDE Sales Report",
        "period": "Period",
        "as_of": "Data as of",
        "note": "Successful simulated payments. Minsk dates, BYN amounts after discounts.",
        "total": "Total",
        "empty": "No paid orders in the selected period.",
        "headers": ["Order", "Paid date", "Items", "Subtotal", "Discount", "Total BYN"],
        "sheet": "Sales",
        "source": "Source: NORTHSIDE PostgreSQL, orders and successful payments.",
    },
}


def content(data, locale):
    words = TEXT[locale]
    rows = [
        [
            str(r["id"]),
            r["paid_at"].astimezone(MINSK).strftime("%Y-%m-%d %H:%M"),
            str(r["item_count"]),
            r["subtotal_byn"],
            r["discount_byn"],
            r["total_byn"],
        ]
        for r in data["rows"]
    ]
    total = [
        words["total"],
        "",
        str(sum(r["item_count"] for r in data["rows"])),
        data["totals"]["subtotal_byn"],
        data["totals"]["discount_byn"],
        data["totals"]["total_byn"],
    ]
    period = f"{words['period']}: {data['start']} - {data['end']}"
    as_of = f"{words['as_of']}: {data['as_of'].astimezone(MINSK):%Y-%m-%d %H:%M:%S}"
    return words, rows, total, period, as_of


def docx_export(data, locale):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    words, rows, total, period, as_of = content(data, locale)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph(words["title"], "Title")
    doc.add_paragraph(period + "\n" + as_of)
    doc.add_paragraph(words["note"])
    table = doc.add_table(rows=1, cols=6)
    table.autofit = False
    widths = [2.1, 4.3, 1.7, 3.3, 3.3, 3.3]
    for cell, width, label in zip(table.rows[0].cells, widths, words["headers"], strict=True):
        cell.text, cell.width = label, Cm(width)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in [*rows, total]:
        for cell, value, width in zip(table.add_row().cells, values, widths, strict=True):
            cell.text, cell.width = str(value), Cm(width)
    for index, row in enumerate(table.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for col, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement("w:" + side)
                for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    edge.set(qn("w:" + key), val)
                borders.append(edge)
            props.append(borders)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCE6EF" if index == 0 else "FFFFFF")
            props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(5)
                paragraph.paragraph_format.space_after = Pt(5)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.RIGHT if col >= 2 else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    run.bold = index in (0, len(table.rows) - 1)
    if not rows:
        doc.add_paragraph(words["empty"])
    doc.add_paragraph(words["source"])
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def pdf_export(data, locale):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    words, rows, total, period, as_of = content(data, locale)
    font = Path(__file__).parent / "assets" / "manrope-regular.ttf"
    if "Northside" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("Northside", str(font)))
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=40,
        bottomMargin=40,
        title=words["title"],
    )
    body = ParagraphStyle("body", fontName="Northside", fontSize=10, leading=15, spaceAfter=9)
    title = ParagraphStyle("title", parent=body, fontSize=19, leading=25, spaceAfter=16)
    small = ParagraphStyle("cell", parent=body, fontSize=9, leading=13, spaceAfter=0)
    table_rows = [
        [Paragraph(escape(str(v)), small) for v in row] for row in [words["headers"], *rows, total]
    ]
    table = Table(table_rows, colWidths=[55, 133, 47, 96, 96, 96], repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6EF")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F3F5F7")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story = [
        Paragraph(escape(words["title"]), title),
        Paragraph(escape(period), body),
        Paragraph(escape(as_of), body),
        Paragraph(escape(words["note"]), body),
        table,
        Spacer(1, 14),
        Paragraph(escape(words["source"]), body),
    ]
    if not rows:
        story.append(Paragraph(escape(words["empty"]), body))

    def footer(canvas, _doc):
        canvas.setFont("Northside", 9)
        canvas.drawRightString(A4[0] - 36, 22, str(_doc.page))

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def xlsx_export(data, locale):
    """A small OOXML worksheet with typed dates/numbers and a cached SUM total."""
    words, rows, total, period, as_of = content(data, locale)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    sheet_rows = []

    def text_cell(ref, value, style=0):
        return f'<c r="{ref}" s="{style}" t="inlineStr"><is><t>{escape(str(value))}</t></is></c>'

    def number_cell(ref, value, style=2):
        return f'<c r="{ref}" s="{style}"><v>{value}</v></c>'

    for index, value in enumerate([words["title"], period, as_of, words["source"]], 1):
        sheet_rows.append(
            f'<row r="{index}" ht="26" customHeight="1">'
            + text_cell(f"A{index}", value, 1 if index == 1 else 0)
            + "</row>"
        )
    sheet_rows.append(
        '<row r="6" ht="30" customHeight="1">'
        + "".join(
            text_cell(f"{chr(65 + i)}6", label, 1) for i, label in enumerate(words["headers"])
        )
        + "</row>"
    )
    for index, source in enumerate(data["rows"], 7):
        local = source["paid_at"].astimezone(MINSK).replace(tzinfo=None)
        excel_date = Decimal(str((local - datetime(1899, 12, 30)).total_seconds())) / 86400
        cells = [
            text_cell(f"A{index}", source["id"]),
            number_cell(f"B{index}", excel_date, 3),
            number_cell(f"C{index}", source["item_count"], 0),
        ]
        for col, key in zip("DEF", ("subtotal_byn", "discount_byn", "total_byn"), strict=True):
            cells.append(number_cell(f"{col}{index}", source[key]))
        sheet_rows.append(f'<row r="{index}" ht="23" customHeight="1">' + "".join(cells) + "</row>")
    last = 7 + len(rows)
    cells = [text_cell(f"A{last}", words["total"], 1)]
    for col, value in zip("CDEF", total[2:], strict=True):
        formula = f"SUM({col}7:{col}{last - 1})" if rows else "0"
        cells.append(
            f'<c r="{col}{last}" s="{0 if col == "C" else 2}"><f>{formula}</f><v>{value}</v></c>'
        )
    sheet_rows.append(f'<row r="{last}" ht="26" customHeight="1">' + "".join(cells) + "</row>")
    worksheet = f'''<worksheet xmlns="{ns}"><sheetViews><sheetView workbookViewId="0">
        <pane ySplit="6" topLeftCell="A7" activePane="bottomLeft" state="frozen"/>
        </sheetView></sheetViews><cols>
        <col min="1" max="1" width="16" customWidth="1"/>
        <col min="2" max="2" width="24" customWidth="1"/>
        <col min="3" max="3" width="12" customWidth="1"/>
        <col min="4" max="6" width="18" customWidth="1"/></cols>
        <sheetData>{"".join(sheet_rows)}</sheetData>
        <autoFilter ref="A6:F{max(6, last - 1)}"/>
        <mergeCells count="4">{"".join(f'<mergeCell ref="A{i}:F{i}"/>' for i in range(1, 5))}
        </mergeCells><pageMargins left="0.4" right="0.4" top="0.5" bottom="0.5"
        header="0.2" footer="0.2"/><pageSetup paperSize="9" orientation="landscape"/>
        </worksheet>'''
    styles = f'''<styleSheet xmlns="{ns}"><numFmts count="1">
        <numFmt numFmtId="164" formatCode="yyyy-mm-dd hh:mm"/></numFmts>
        <fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>
        <font><b/><sz val="11"/><name val="Calibri"/></font></fonts>
        <fills count="3"><fill><patternFill patternType="none"/></fill>
        <fill><patternFill patternType="gray125"/></fill>
        <fill><patternFill patternType="solid"><fgColor rgb="FFDCE6EF"/>
        <bgColor indexed="64"/></patternFill></fill></fills>
        <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
        <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
        <cellXfs count="4"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
        <xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFill="1"/>
        <xf numFmtId="4" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>
        <xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>
        </cellXfs><cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/>
        </cellStyles></styleSheet>'''
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    files = {
        "[Content_Types].xml": """<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
            <Default Extension="rels"
                ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
            <Default Extension="xml" ContentType="application/xml"/>
            <Override PartName="/xl/workbook.xml"
                ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
            <Override PartName="/xl/worksheets/sheet1.xml"
                ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
            <Override PartName="/xl/styles.xml"
                ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
            </Types>""",
        "_rels/.rels": f'<Relationships xmlns="{rel_ns}"><Relationship Id="rId1" '
        f'Type="{office_rel}officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{ns}" xmlns:r="{office_rel[:-1]}">'
        f'<sheets><sheet name="{words["sheet"]}" sheetId="1" r:id="rId1"/>'
        '</sheets><calcPr calcId="191029" fullCalcOnLoad="1"/></workbook>',
        "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{rel_ns}"><Relationship Id="rId1" '
        f'Type="{office_rel}worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" '
        f'Type="{office_rel}styles" Target="styles.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": worksheet,
        "xl/styles.xml": styles,
    }
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, xml in files.items():
            archive.writestr(name, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>' + xml)
    return output.getvalue()


EXPORTERS = {"docx": docx_export, "pdf": pdf_export, "xlsx": xlsx_export}
